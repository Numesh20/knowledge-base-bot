# -*- coding: utf-8 -*-
"""
Knowledge Base Bot  -  bot.py
Uses Google Gemini (FREE) + ChromaDB (local)
Run:  py bot.py
"""

import os
import sys
import glob
from dotenv import load_dotenv

# Fix Windows terminal encoding
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stdin.reconfigure(encoding="utf-8", errors="replace")

# ── Load API key from .env file ──
load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

if not GEMINI_API_KEY or GEMINI_API_KEY == "your-gemini-api-key-here":
    print("\n[ERROR] No GEMINI_API_KEY found in your .env file!")
    print("   1. Open the .env file")
    print("   2. Replace 'your-gemini-api-key-here' with your real key")
    print("   3. Get a FREE key at: https://aistudio.google.com\n")
    exit(1)

print("[OK] API key loaded.")

# ── Import AI libraries ──
print("[...] Loading AI libraries...")
import chromadb
from google import genai
from google.genai import types

# ── Configure Gemini client ──
client = genai.Client(api_key=GEMINI_API_KEY)

# ── Helper: extract text from file (.txt, .pdf, .docx, .csv, .md) ──
def extract_text_from_file(file_path):
    """Extract text from supported files: .txt, .pdf, .docx, .csv, .md."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    elif ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                pages_text.append(t.strip())
        return "\n\n".join(pages_text)
    elif ext == ".docx":
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            tables_text = []
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        tables_text.append(" | ".join(row_data))
            all_content = paragraphs + tables_text
            return "\n\n".join(all_content)
        except Exception as e:
            print(f"  [ERROR] Failed to read Word document {file_path}: {e}")
            return ""
    elif ext == ".csv":
        try:
            import csv
            rows_text = []
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                header = None
                for i, row in enumerate(reader):
                    clean_row = [cell.strip() for cell in row if cell.strip()]
                    if not clean_row:
                        continue
                    if header is None:
                        header = clean_row
                        rows_text.append("Columns: " + ", ".join(header))
                    else:
                        row_items = []
                        for col_idx, val in enumerate(clean_row):
                            col_name = header[col_idx] if col_idx < len(header) else f"Column {col_idx+1}"
                            row_items.append(f"{col_name}: {val}")
                        rows_text.append(f"Row {i}: " + "; ".join(row_items))
            return "\n".join(rows_text)
        except Exception as e:
            print(f"  [ERROR] Failed to read CSV file {file_path}: {e}")
            return ""
    return ""

# ── Helper: split text into chunks ──
def chunk_text(text, chunk_size=400, overlap=50):
    """Break a big text into smaller overlapping pieces."""
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i : i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        i += chunk_size - overlap
    return chunks

# ── Helper: get embedding from Gemini ──
def get_embedding(text, max_retries=3):
    """Convert text into a list of numbers (vector) using Gemini with retry logic."""
    import time
    for attempt in range(max_retries):
        try:
            response = client.models.embed_content(
                model="gemini-embedding-001",
                contents=text,
            )
            return response.embeddings[0].values
        except Exception as e:
            if attempt == max_retries - 1:
                raise e
            print(f"      [RETRY] Embedding call failed ({e}). Retrying in {1.5 * (attempt + 1):.1f}s...")
            time.sleep(1.5 * (attempt + 1))

# ── Build the Knowledge Base ──
def build_knowledge_base():
    """Load all supported documents from /knowledge folder and store in ChromaDB."""
    print("\n[...] Building knowledge base...")

    # Connect to local ChromaDB (saves data in ./chroma_db folder)
    db = chromadb.PersistentClient(path="./chroma_db")

    # Delete old collection if it exists (fresh rebuild)
    try:
        db.delete_collection("knowledge_base")
    except Exception:
        pass

    collection = db.create_collection("knowledge_base")

    # Find all supported files in the knowledge/ folder (.txt, .pdf, .docx, .csv, .md)
    files = []
    for ext in [".txt", ".pdf", ".docx", ".csv", ".md"]:
        files.extend(glob.glob(f"knowledge/*{ext}"))
    files = sorted(files)

    if not files:
        print("[WARN] No supported files found in the knowledge/ folder!")
        print("       Add some .txt, .pdf, .docx, .csv, or .md files there and run again.")
        return collection, 0

    total_chunks = 0
    for file_path in files:
        print(f"   [FILE] Loading: {file_path}")
        try:
            text = extract_text_from_file(file_path)
        except Exception as e:
            print(f"   [ERROR] Could not read {file_path}: {e}")
            continue

        if not text.strip():
            print(f"   [WARN] No readable text found in {file_path}")
            continue

        # Split into chunks
        chunks = chunk_text(text)
        print(f"          -> Split into {len(chunks)} chunks")

        # Store each chunk with its embedding
        for i, chunk in enumerate(chunks):
            embedding = get_embedding(chunk)
            collection.add(
                ids=[f"{file_path}-chunk-{i}"],
                embeddings=[embedding],
                documents=[chunk],
                metadatas=[{"source": file_path, "chunk": i}],
            )
        total_chunks += len(chunks)

    print(f"\n[READY] Knowledge base built! ({len(files)} files, {total_chunks} chunks)\n")
    return collection, total_chunks

# ── Delete document from disk and ChromaDB ──
def delete_document(filename, db_path="./chroma_db"):
    """Delete a document from knowledge/ and purge its chunks from ChromaDB."""
    safe_filename = os.path.basename(filename)
    file_path = os.path.join("knowledge", safe_filename)

    file_existed = False
    if os.path.exists(file_path):
        os.remove(file_path)
        file_existed = True

    db = chromadb.PersistentClient(path=db_path)
    try:
        col = db.get_collection("knowledge_base")
        all_data = col.get()
        ids_to_delete = []
        for doc_id, meta in zip(all_data.get("ids", []), all_data.get("metadatas", [])):
            src = meta.get("source", "") if meta else ""
            if os.path.basename(src) == safe_filename or safe_filename in doc_id:
                ids_to_delete.append(doc_id)
        if ids_to_delete:
            col.delete(ids=ids_to_delete)
        deleted_count = len(ids_to_delete)
    except Exception as e:
        deleted_count = 0
        print(f"[WARN] ChromaDB error during deletion: {e}")

    print(f"[DELETED] '{safe_filename}' (file removed: {file_existed}, {deleted_count} chunks purged)")
    return file_existed or deleted_count > 0

# ── Retrieve relevant chunks ──
def retrieve(collection, question, top_k=3):
    """Find the most relevant chunks for a question."""
    question_embedding = get_embedding(question)
    results = collection.query(
        query_embeddings=[question_embedding],
        n_results=top_k,
    )
    chunks = results["documents"][0]
    sources = [m["source"] for m in results["metadatas"][0]]
    return chunks, sources

# ── Generate answer with Gemini ──
def generate_answer(question, context_chunks):
    """Send question + retrieved context to Gemini for an answer."""
    context = "\n\n".join(context_chunks)

    prompt = f"""You are a helpful knowledge base assistant.
Answer the user's question using ONLY the information provided below.
If the answer is not in the context, say: "I don't have that information in my knowledge base."

=== KNOWLEDGE BASE CONTEXT ===
{context}
==============================

User Question: {question}

Answer:"""

    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents=prompt,
    )
    return response.text.strip()

# ── Main Chat Loop ──
def main():
    print("=" * 55)
    print("  Knowledge Base Bot  (powered by Google Gemini)")
    print("=" * 55)

    # Build the knowledge base from your files
    collection, total_chunks = build_knowledge_base()

    if total_chunks == 0:
        print("No knowledge loaded. Add .txt files to knowledge/ folder.")
        return

    print("Ask me anything! (type 'quit' to exit)\n")

    chat_history = []

    while True:
        question = input("You: ").strip()

        if not question:
            continue
        if question.lower() in ("quit", "exit", "bye"):
            print("\nBot: Goodbye!\n")
            break

        # Step 1: Retrieve relevant chunks
        print("[...] Searching knowledge base...")
        chunks, sources = retrieve(collection, question)

        # Step 2: Generate answer
        print("[...] Generating answer...\n")
        answer = generate_answer(question, chunks)

        # Step 3: Show answer + sources
        print(f"Bot: {answer}")
        print(f"\nSources: {', '.join(set(sources))}")
        print("-" * 55)

        # Save to history
        chat_history.append({"question": question, "answer": answer})

if __name__ == "__main__":
    main()
